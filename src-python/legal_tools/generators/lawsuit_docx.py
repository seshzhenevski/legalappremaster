# -*- coding: utf-8 -*-
"""
generators/lawsuit_docx.py
──────────────────────────
Генерация искового заявления и описи (python-docx).
"""
from __future__ import annotations

import io
from datetime import date, timedelta
from decimal import Decimal
from typing import List, Optional, Tuple

from ..config import (
    LAW_PLAINTIFF_NAME, LAW_PLAINTIFF_INN, LAW_PLAINTIFF_OGRN, LAW_PLAINTIFF_ADDRESS,
    LAW_PLAINTIFF_EMAIL, LAW_COURT_HEADER, LAW_COURT_ADDRESS, LAW_REPRESENTATIVE,
)
from ..core.formatting import fmt, normalize_quotes, strip_all_quotes, parse_contract_ref, next_day
from ..core.duty import plural_form
from ..assets.embedded import get_opis_template_bytes

try:
    import docx as _docxlib
    from docx import Document as DocxDocument
    from docx.shared import Pt as DocxPt, Cm as DocxCm, RGBColor as DocxRGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn as _docx_qn
    from docx.oxml import OxmlElement as _docx_oxml_element
    _DOCX_OK = True
except Exception:
    _DOCX_OK = False

_opis_cache = None


def money_words(amount: Decimal) -> str:
    """'1 263 316,94 рублей' — сумма с верным словом «рубль/рубля/рублей»."""
    rubles_int = int(amount)
    word = plural_form(rubles_int, "рубль", "рубля", "рублей")
    return f"{fmt(amount)} {word}"


def money_whole(amount: Decimal) -> str:
    """'62 900 рублей' — целая сумма без копеек (для госпошлины)."""
    rubles_int = int(amount)
    word = plural_form(rubles_int, "рубль", "рубля", "рублей")
    formatted = f"{rubles_int:,}".replace(",", " ")
    return f"{formatted} {word}"




def _law_p_run(paragraph, text, bold=False, size=11):
    r = paragraph.add_run(text)
    r.font.name = "Calibri Light"
    r.font.size = DocxPt(size)
    r.bold = bold
    return r


def _law_set_table_borders(table, color="999999", sz=4):
    tblPr = table._tbl.tblPr
    borders = _docx_oxml_element('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = _docx_oxml_element(f'w:{edge}')
        el.set(_docx_qn('w:val'), 'single')
        el.set(_docx_qn('w:sz'), str(sz))
        el.set(_docx_qn('w:space'), '0')
        el.set(_docx_qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def _law_cell(cell, text="", bold=False, align=None, size=11, shade=None, nowrap=True):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    if text:
        _law_p_run(p, text, bold=bold, size=size)
    tcPr = cell._tc.get_or_add_tcPr()
    if nowrap:
        tcPr.append(_docx_oxml_element('w:noWrap'))
    if shade:
        shd = _docx_oxml_element('w:shd')
        shd.set(_docx_qn('w:val'), 'clear')
        shd.set(_docx_qn('w:color'), 'auto')
        shd.set(_docx_qn('w:fill'), shade)
        tcPr.append(shd)
    return cell


def _law_set_widths(table, widths_pt):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = _docx_oxml_element('w:tblLayout')
    layout.set(_docx_qn('w:type'), 'fixed')
    tblPr.append(layout)
    for row in table.rows:
        for idx, w in enumerate(widths_pt):
            if idx < len(row.cells):
                row.cells[idx].width = DocxPt(w)
    grid = table._tbl.tblGrid
    cols = grid.findall(_docx_qn('w:gridCol'))
    for idx, w in enumerate(widths_pt):
        if idx < len(cols):
            cols[idx].set(_docx_qn('w:w'), str(int(w * 20)))


def _law_cm_to_dxa(cm: float) -> int:
    return int(round(cm / 2.54 * 72 * 20))


def _law_set_cell_margins(table, top_cm=0.05, bottom_cm=0.05, left_cm=0.1, right_cm=0.1):
    top = _law_cm_to_dxa(top_cm)
    bottom = _law_cm_to_dxa(bottom_cm)
    left = _law_cm_to_dxa(left_cm)
    right = _law_cm_to_dxa(right_cm)
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            mar = _docx_oxml_element('w:tcMar')
            for edge, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
                el = _docx_oxml_element(f'w:{edge}')
                el.set(_docx_qn('w:w'), str(val))
                el.set(_docx_qn('w:type'), 'dxa')
                mar.append(el)
            tcPr.append(mar)


def _law_build_invoice_table(doc, table_rows: List[dict]):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, htext in enumerate(["Дата счета", "№ счета", "Сумма долга", "Срок оплаты"]):
        _law_cell(hdr[i], htext, align=WD_ALIGN_PARAGRAPH.CENTER, shade="BDD6EE")

    for row in table_rows:
        cells = table.add_row().cells
        _law_cell(cells[0], row["invoice_date"].strftime("%d.%m.%Y"), align=WD_ALIGN_PARAGRAPH.CENTER)
        _law_cell(cells[1], row["number"], align=WD_ALIGN_PARAGRAPH.CENTER)
        _law_cell(cells[2], fmt(row["amount"]), align=WD_ALIGN_PARAGRAPH.RIGHT)
        _law_cell(cells[3], row["due_date"].strftime("%d.%m.%Y"), align=WD_ALIGN_PARAGRAPH.CENTER)

    _law_set_widths(table, [120, 127, 134, 120])
    _law_set_table_borders(table, color="A0A0A4", sz=4)
    return table


def _law_build_penalty_table(
    doc, blocks: List[dict], total_debt: Decimal, total_penalty: Decimal,
    cap_info: Optional[Tuple[Decimal, Decimal, Decimal]] = None,
):
    table = doc.add_table(rows=2, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for col, text in enumerate(["Месяц", "Начислено", "Долг"]):
        merged = table.cell(0, col).merge(table.cell(1, col))
        _law_cell(merged, text, align=WD_ALIGN_PARAGRAPH.RIGHT)

    period_cell = table.cell(0, 3).merge(table.cell(0, 5))
    _law_cell(period_cell, "Период просрочки", align=WD_ALIGN_PARAGRAPH.CENTER)
    for col, text in zip((3, 4, 5), ("с", "по", "дней")):
        _law_cell(table.cell(1, col), text, align=WD_ALIGN_PARAGRAPH.CENTER)

    for col, text in zip((6, 7), ("Формула", "Пени")):
        merged = table.cell(0, col).merge(table.cell(1, col))
        _law_cell(merged, text, align=WD_ALIGN_PARAGRAPH.RIGHT)

    for block in blocks:
        rows = block["rows"]
        n = len(rows)
        start_row_idx = len(table.rows)
        for sub in rows:
            cells = table.add_row().cells
            if sub[0] == "debt":
                _, debt_s, c_s, po_s, days_s, formula_s, peni_s = sub
                for col, val in zip(range(2, 8), (debt_s, c_s, po_s, days_s, formula_s, peni_s)):
                    _law_cell(cells[col], val, align=WD_ALIGN_PARAGRAPH.RIGHT)
            else:
                _, debt_s, date_s = sub
                _law_cell(cells[2], debt_s, align=WD_ALIGN_PARAGRAPH.RIGHT)
                _law_cell(cells[3], date_s, align=WD_ALIGN_PARAGRAPH.RIGHT)
                merged = cells[4].merge(cells[6])
                _law_cell(merged, "Погашение части долга", align=WD_ALIGN_PARAGRAPH.LEFT)
                _law_cell(cells[7], "")

        if n > 1:
            m_cell = table.cell(start_row_idx, 0).merge(table.cell(start_row_idx + n - 1, 0))
            n_cell = table.cell(start_row_idx, 1).merge(table.cell(start_row_idx + n - 1, 1))
        else:
            m_cell = table.cell(start_row_idx, 0)
            n_cell = table.cell(start_row_idx, 1)
        _law_cell(m_cell, block["month"], align=WD_ALIGN_PARAGRAPH.RIGHT)
        _law_cell(n_cell, block["nach"], align=WD_ALIGN_PARAGRAPH.RIGHT)

        itogo_cells = table.add_row().cells
        merged_label = itogo_cells[0].merge(itogo_cells[5])
        _law_cell(merged_label, "")
        _law_cell(itogo_cells[6], "Итого:", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _law_cell(itogo_cells[7], block["itogo"] + " руб.", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    for label in (
        f"Сумма основного долга: {fmt(total_debt)} руб.",
        f"Сумма пеней по всем задолженностям: {fmt(total_penalty)} руб.",
    ):
        cells = table.add_row().cells
        merged = cells[0]
        for c in cells[1:]:
            merged = merged.merge(c)
        _law_cell(merged, label, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    if cap_info:
        cap_pct_v, cap_value, raw_value = cap_info
        note = (
            f"Применено ограничение по договору: не более {cap_pct_v}% от суммы долга — "
            f"неустойка уменьшена с {fmt(raw_value)} руб. до {fmt(cap_value)} руб."
        )
        cells = table.add_row().cells
        merged = cells[0]
        for c in cells[1:]:
            merged = merged.merge(c)
        _law_cell(merged, note, align=WD_ALIGN_PARAGRAPH.LEFT)

    _law_set_table_borders(table, color="CCCCCC", sz=4)
    _law_set_widths(table, [58, 57, 62, 56, 57, 29, 120, 86])
    _law_set_cell_margins(table, top_cm=0.05, bottom_cm=0.05, left_cm=0.1, right_cm=0.1)
    return table


def generate_lawsuit_docx(path: str, ctx: dict) -> None:
    """Строит docx искового заявления «с нуля» (python-docx), полностью
    повторяя структуру/оформление образца Исковое_заявление.docx."""
    if not _DOCX_OK:
        raise RuntimeError("Библиотека python-docx не установлена.")

    d = ctx["defendant"]
    doc = DocxDocument()

    section = doc.sections[0]
    section.page_width = DocxCm(21.0)
    section.page_height = DocxCm(29.7)
    section.top_margin = DocxCm(1.25)
    section.bottom_margin = DocxCm(1.25)
    section.left_margin = DocxCm(1.75)
    section.right_margin = DocxCm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri Light"
    normal.font.size = DocxPt(11)
    normal.paragraph_format.space_before = DocxPt(0)
    normal.paragraph_format.space_after = DocxPt(0)

    def add_par(text="", bold=False, align=None, justify_body=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = DocxPt(0)
        p.paragraph_format.space_after = DocxPt(0)
        if justify_body:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = DocxCm(1.25)
        elif align is not None:
            p.alignment = align
        if text:
            _law_p_run(p, text, bold=bold)
        return p

    add_par(LAW_COURT_HEADER, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par(LAW_COURT_ADDRESS, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par("", align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _law_p_run(p, "Истец:", bold=True); _law_p_run(p, f" {LAW_PLAINTIFF_NAME}")
    add_par(f"ИНН {LAW_PLAINTIFF_INN}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par(f"ОГРН {LAW_PLAINTIFF_OGRN}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par(f"Адрес: {LAW_PLAINTIFF_ADDRESS}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par(f"e-mail: {LAW_PLAINTIFF_EMAIL}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par("", align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _law_p_run(p, "Ответчик:", bold=True); _law_p_run(p, f" {d['name']}")
    add_par(f"ИНН {d['inn']}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par(f"ОГРН {d['ogrn']}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par(f"Адрес: {d['address']}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_par("", align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _law_p_run(p, "Цена иска:", bold=True)
    _law_p_run(p, f" {money_words(ctx['claim_amount'])}.")
    add_par("", align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _law_p_run(p, "Исковое заявление", bold=True)
    add_par("")

    contract_full = ctx["contract_ref"]

    add_par(
        f"Между {LAW_PLAINTIFF_NAME} (далее – Исполнитель) и {d['name']} "
        f"(далее – Заказчик) был заключен договор № {contract_full} (далее - Договор).",
        justify_body=True,
    )
    add_par(
        "Исходя из предмета заключенного Договора, Исполнитель предоставлял "
        "Заказчику услуги по организации деловых и иных поездок в соответствии "
        "с заказами Заказчика.",
        justify_body=True,
    )
    add_par(
        "Согласно п. 1.2. Услуги по Договору оказываются с использованием ИТ "
        "платформы – онлайн системы управления поездками «Trivio» "
        "https://trivio.ru/ (далее – Сервис).",
        justify_body=True,
    )
    add_par(
        "Цена заказа состояла из стоимости услуг Исполнителя (сервисные сборы) "
        "и затрат Исполнителя на выполнение заказа Заказчика, включая стоимость "
        "всех включенных в заказ услуг (то есть фактическая стоимость "
        "авиабилетов, проживания, трансферов и т.д.).",
        justify_body=True,
    )
    add_par(
        "В силу пункта 3.8. Договора Заказчик обязан произвести оплату "
        "Исполнителю за оказанные услуги, в соответствии с выставленными "
        "Исполнителем счетами, в течение 14 календарных дней с момента "
        "получения счета от Исполнителя.",
        justify_body=True,
    )
    add_par(
        "Исполнитель оказал услуги надлежащего качества в соответствии с "
        "требованиями Договора и в полном объеме выполнил свои обязательства "
        "перед Заказчиком, на основании чего Исполнитель выставил Заказчику "
        "счета на оплату:",
        justify_body=True,
    )

    _law_build_invoice_table(doc, ctx["table_rows"])

    add_par("")
    p = add_par(justify_body=True)
    _law_p_run(p, "Итоговая задолженность Заказчика составляет ")
    _law_p_run(p, money_words(ctx["total_debt"]), bold=True)
    _law_p_run(p, ".")
    add_par("")

    edi_date_suffix = f" от {ctx['contract_date']}" if ctx.get("contract_date") else ""
    add_par(
        "В рамках заключенного Договора между сторонами было заключено "
        f"соглашение о переходе на электронный документооборот{edi_date_suffix} "
        "(далее – Соглашение), которое изложено в Приложении № 2 к Договору "
        "и является его неотъемлемой частью в силу п. 9.4. Договора.",
        justify_body=True,
    )
    if ctx.get("docs_signed", True):
        # Стандартный вариант: закрывающие документы подписаны обеими сторонами
        add_par(
            "Согласно пункту 3.7. Договора Исполнитель по средствам согласованного "
            "между сторонами электронного документооборота предоставил Заказчику "
            "первичные бухгалтерские документы, предусмотренные условиями договора. "
            "Услуги приняты Заказчиком. Универсальные передаточные документы и Акты "
            "приема-передачи авиа и жд билетов по указанным выше счетам подписаны "
            "Ответчиком посредством ЭДО. Таким образом факт оказания услуг "
            "надлежащим образом подтвержден.",
            justify_body=True,
        )
    else:
        # Альтернативный вариант: часть закрывающих документов не подписана
        _unsigned_paras = [
            (
                "Исходя из условий указанного Соглашения Стороны согласовали обмен "
                "формализованными и неформализованными электронными документами (в том числе "
                "УПД, акты-приема передачи билетов и др.), подписанными усиленной "
                "квалифицированной электронной подписью, в рамках электронного документооборота."
            ),
            (
                "Согласно п. 5.1., 5.3. Соглашения, любая из Сторон может в любой момент "
                "отказаться от участия в электронном документообороте, направив уведомление "
                "об этом другой Стороне в системе ЭДО или на бумажном носителе, за 30 "
                "(Тридцать) календарных дней до прекращения использования электронного "
                "документооборота и прекращения действия Соглашения. Также Стороны обязаны "
                "информировать друг друга о невозможности обмена документами в электронном виде."
            ),
            (
                "Подобных уведомлений в адрес Исполнителя не поступало, равно как и информации "
                "о невозможности обмена документами в электронном виде со стороны Заказчика, "
                "в связи с чем Исполнитель разумно и добросовестно направлял закрывающие "
                "документы Заказчику посредством согласованной сторонами системы электронного "
                "документооборота."
            ),
            (
                "Согласно пункту 3.6. Договора Исполнитель по средствам согласованного между "
                "сторонами электронного документооборота предоставил Заказчику первичные "
                "бухгалтерские документы, предусмотренные условиями договора. "
            ),
            (
                "В качестве подтверждения факта оказания услуг Исполнитель ссылается на акты "
                "приема-передачи авиа и ж/д билетов и Универсальные передаточные документы "
                "(УПД), которые объединяют в себе акт выполненных работ/оказанных услуг и "
                "счет-фактуру, и содержат всю необходимую информацию об оказанных услугах, "
                "датах их оказания и стоимости."
            ),
            (
                "В силу положений закона о свободе договора и её пределах, стороны вправе "
                "самостоятельно определить порядок документооборота. В рамках настоящего "
                "Договора сторонами был определен следующий порядок:"
            ),
            (
                "В силу п. 3.6. Договора, Заказчик обязан возвратить подписанные "
                "уполномоченным представителем Заказчика первичные бухгалтерские документы "
                "не позднее 2 рабочих дней с момента получения. "
            ),
            (
                "В случае не подписания и(или) невозврата подписанных со стороны Заказчика "
                "экземпляров первичных бухгалтерских документов в течение 2 рабочих дней с "
                "момента их получения (и/или не представления мотивированного письменного "
                "отказа от их подписания в этот же срок), односторонне подписанные "
                "Исполнителем первичные бухгалтерские документы, предусмотренные Договором, "
                "считаются подтверждением надлежащего, в полном объеме, оказания услуг по "
                "Договору."
            ),
            (
                "Факт надлежащей отправки Исполнителем первичных бухгалтерских документов "
                "подтверждается соответствующим штампом оператора электронного документооборота "
                "СКБ Контур.Диадок."
            ),
            (
                "Услуги были приняты Заказчиком. Универсальные передаточные документы и Акты "
                "приема-передачи авиа и жд билетов подписаны Ответчиком посредством ЭДО. "
                "Вместе с тем, по ряду закрывающих документов в установленный Договором срок "
                "Заказчик свою обязанность по предоставлению Исполнителю подписанных со своей "
                "стороны первичных бухгалтерских документов не выполнил, мотивированного "
                "письменного отказа от их подписания также не предоставил. Доказательства "
                "ненадлежащего оказания услуг либо обстоятельств, исключающих обязанность их "
                "оплаты Заказчиком также не представлено."
            ),
            (
                "Учитывая изложенные обстоятельства, руководствуясь ст. 421 ГК РФ отражающей "
                "принцип свободы договора, факт надлежащего оказания услуг в данном случае "
                "считается полностью подтвержденным. Данный вывод коррелирует с позицией ВС РФ, "
                "который неоднократно указывал, что неподписание акта (УПД) при фактическом "
                "оказании услуг не освобождает Заказчика от обязанности по оплате данных услуг."
            ),
        ]
        for text in _unsigned_paras:
            add_par(text, justify_body=True)
    add_par(
        "В силу ст. 309 ГК РФ обязательства должны исполняться надлежащим "
        "образом, а согласно ст. 310 ГК РФ односторонний отказ от исполнения "
        "обязательства не допускается. Вместе с тем, Заказчик, приняв оказанные "
        "услуги, оплату в установленный Договором срок и до настоящего времени "
        "не произвёл.",
        justify_body=True,
    )
    add_par("")

    rate_str = str(ctx["rate"])
    if "." in rate_str:
        rate_str = rate_str.rstrip("0").rstrip(".")
    rate_display = rate_str.replace(".", ",")
    add_par(
        f"Пунктом 4.7. Договора предусмотрена договорная неустойка в случае "
        f"просрочки Заказчиком срока оплаты, в размере {rate_display}% от "
        f"неоплаченной суммы задолженности за каждый день просрочки оплаты.",
        justify_body=True,
    )
    add_par("Расчет задолженности:", justify_body=True)

    _law_build_penalty_table(doc, ctx["blocks"], ctx["total_debt"], ctx["total_penalty"], ctx.get("cap_info"))

    add_par("")
    claim_date_str = ctx["claim_date"].strftime("%d.%m.%Y")
    p = add_par(justify_body=True)
    _law_p_run(p, f"На {claim_date_str} договорная неустойка в связи с просрочкой "
                  f"оплаты услуг Исполнителя по Договору составила ")
    _law_p_run(p, money_words(ctx["total_penalty"]))
    _law_p_run(p, ".")
    add_par("")

    p = add_par(justify_body=True)
    _law_p_run(p, f"Общая сумма задолженности {d['name']} перед {LAW_PLAINTIFF_NAME} "
                  f"по обязательствам в рамках Договора № {contract_full} составляет ")
    _law_p_run(p, money_words(ctx['claim_amount']), bold=True)
    _law_p_run(p, ", из которых: ")

    p = add_par(justify_body=True)
    _law_p_run(p, "- задолженность по оплате услуг Истца в размере ")
    _law_p_run(p, money_words(ctx['total_debt']), bold=True)
    _law_p_run(p, ", ")

    period_start_str = ctx["period_start"].strftime("%d.%m.%Y")
    p = add_par(justify_body=True)
    _law_p_run(p, f"- договорная неустойка за период с {period_start_str} по "
                  f"{claim_date_str} в соответствии с п. 4.7. Договора в размере ")
    _law_p_run(p, money_words(ctx['total_penalty']), bold=True)
    _law_p_run(p, ".")
    add_par("")

    pretenzia_date_str = ctx["pretenzia_date"].strftime("%d.%m.%Y")
    add_par(
        f"Заказчику была направлена претензия (исх. № {ctx['pretenzia_number']} "
        f"от {pretenzia_date_str}) с целью урегулировать спор в досудебном "
        f"порядке, однако до настоящего времени денежные средства в счет "
        f"оплаты образовавшейся задолженности в адрес Истца не поступали.",
        justify_body=True,
    )
    add_par("")
    add_par(
        "Пунктом 7.3. Договора предусмотрена договорная подсудность, в "
        "соответствии с которой спор подлежит рассмотрению в Арбитражном суде "
        "города Москвы.",
        justify_body=True,
    )
    add_par("В соответствии с изложенным,", justify_body=True)
    add_par("")

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _law_p_run(p, "ПРОШУ:", bold=True)
    add_par("")

    next_day_str = next_day(ctx["claim_date"], False).strftime("%d.%m.%Y")
    demands = [
        f"Взыскать с Ответчика сумму задолженности по оплате услуг Истца по "
        f"договору № {contract_full} в размере {money_words(ctx['total_debt'])}.",

        f"Взыскать с Ответчика в пользу Истца договорную неустойку за период "
        f"с {period_start_str} по {claim_date_str} в соответствии с п. 4.7. "
        f"Договора в размере {money_words(ctx['total_penalty'])}.",

        f"Взыскать с Ответчика в пользу Истца договорную неустойку в размере, "
        f"предусмотренном п. 4.7. Договора, с {next_day_str} по день "
        f"фактического исполнения Ответчиком решения суда.",

        f"Взыскать с Ответчика почтовые расходы в размере "
        f"{money_words(ctx['postal_costs'])}.",

        f"Взыскать с Ответчика в пользу Истца уплаченную госпошлину в размере "
        f"{money_whole(ctx['duty_amount'])}.",
    ]
    for idx, text in enumerate(demands, 1):
        add_par(f"{idx}. {text}", justify_body=True)

    add_par("")
    p = doc.add_paragraph()
    _law_p_run(p, "Приложения:", bold=True)
    add_par("")

    attachments = [
        "Копия документа, подтверждающего направление копии искового заявления Ответчику;",
        "Документы об оплате госпошлины;",
        f"Копия договора № {contract_full};",
        "Копии актов сверки взаимных расчетов, счетов на оплату с актами "
        "приема-передачи билетов и УПД, подтверждающих оказание услуг на "
        "сумму счетов (документы приложены в хронологическом порядке по "
        "принципу счет на оплату, далее следуют документы, подтверждающие "
        "оказание услуг на сумму счета);",
        "Копия досудебной претензии и документов, подтверждающих направление "
        "досудебной претензии Ответчику;",
        "Копия выписки из ЕГРЮЛ на Истца;",
        "Копия выписки из ЕГРЮЛ на Ответчика;",
        "Копия доверенности представителя;",
        "Копия диплома представителя.",
    ]
    for idx, text in enumerate(attachments, 1):
        add_par(f"{idx}. {text}", justify_body=True)

    add_par("")
    add_par("")
    p = doc.add_paragraph()
    _law_p_run(p, f"Представитель {LAW_PLAINTIFF_NAME} ______________/{LAW_REPRESENTATIVE}/", bold=True)

    doc.save(path)


# ── Генерация docx — Опись почтового вложения (редактирование шаблона) ──

def _law_iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                yield p
            for nested in cell.tables:
                yield from _law_iter_table_paragraphs(nested)


def _law_replace_paragraph_text(paragraph, new_text):
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def generate_opis_docx(
    path: str, defendant_name: str, document_title: str = "Исковое заявление",
) -> Tuple[int, int]:
    """Заполняет встроенный шаблон описи (форма ф.107): подставляет
    наименование ответчика и проставляет «1» в столбец «Кол-во предметов»
    (строка предмета и строка итога), на обоих экземплярах формы.

    document_title задаёт название вложения в описи: по умолчанию «Исковое
    заявление» (для генератора иска), либо, например, «Досудебная претензия»
    для генератора претензий."""
    if not _DOCX_OK:
        raise RuntimeError("Библиотека python-docx не установлена.")

    global _opis_cache
    if _opis_cache is None:
        _opis_cache = get_opis_template_bytes()
    doc = DocxDocument(io.BytesIO(_opis_cache))

    target_giti = "Исковое заявление ООО «ГИТИ»"
    replaced_name = 0
    replaced_count = 0
    for table in doc.tables:
        for p in _law_iter_table_paragraphs(table):
            if p.text == target_giti:
                _law_replace_paragraph_text(p, f"{document_title} {defendant_name}")
                replaced_name += 1
            elif p.text.strip() in ("2 листа", "2"):
                # «2 листа» — ячейка строки предмета; «2» — ячейка итогов.
                # В обоих случаях ставим «1»: в описи всегда одно исковое заявление.
                _law_replace_paragraph_text(p, "1")
                replaced_count += 1

    doc.save(path)
    return replaced_name, replaced_count


# ══════════════════════════════════════════════════════════════════════
# Фоновый поток: Excel → расчёт → приложения (Сортировщик) → госпошлина → docx
# ══════════════════════════════════════════════════════════════════════
