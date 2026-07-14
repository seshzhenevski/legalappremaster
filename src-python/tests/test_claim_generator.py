# -*- coding: utf-8 -*-
"""Юнит-тесты слоя логики: генерация досудебной претензии."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from unittest import mock

import openpyxl

from legal_tools.config import (
    CLAIM_REGISTRY_DEFAULT_FOLDER,
    CLAIM_REGISTRY_FILENAME,
)
from logic.claim_generator import (
    parse_iso_date,
    resolve_registry_path,
    summarize_invoices_excel,
    peek_next_claim_number,
    append_registry_row,
    generate_claim_package,
    export_claim_document,
    find_claim_by_inn,
)
from legal_tools.generators.claim_text import (
    format_claim_number,
    letter_reference,
    money_phrase,
    contract_phrase,
)
from legal_tools.config import CLAIM_REGISTRY_FILENAME


class HelperParsingTests(unittest.TestCase):
    """Проверки вспомогательных преобразований генератора претензии."""

    def test_parse_iso_date_value(self):
        self.assertEqual(parse_iso_date("2026-07-13"), date(2026, 7, 13))

    def test_parse_iso_date_empty_returns_none(self):
        self.assertIsNone(parse_iso_date(""))
        self.assertIsNone(parse_iso_date(None))

    def test_format_claim_number_strips_prefix(self):
        self.assertEqual(format_claim_number("№57"), "57")
        self.assertEqual(format_claim_number(" 57 "), "57")
        self.assertEqual(format_claim_number("57"), "57")
        self.assertEqual(format_claim_number(""), "")

    def test_letter_reference(self):
        self.assertEqual(
            letter_reference("57", date(2026, 7, 13)), "Исх. № 57 от 13.07.2026"
        )

    def test_money_phrase_plural(self):
        self.assertEqual(money_phrase(Decimal("269789.80")), "269 789,80 рублей")
        self.assertEqual(money_phrase(Decimal("1")), "1,00 рубль")

    def test_contract_phrase_variants(self):
        self.assertEqual(
            contract_phrase("2706/23-АР", date(2023, 6, 27)),
            "№ 2706/23-АР от 27.06.2023",
        )
        self.assertEqual(contract_phrase("2706/23-АР", None), "№ 2706/23-АР")
        self.assertEqual(contract_phrase("", None), "№ ______")


class RegistryFileTests(unittest.TestCase):
    """Проверки работы с реестром «ОТПРАВКИ ПРЕТЕНЗИЙ.xlsx»."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.registry_path = Path(self.temp_dir) / CLAIM_REGISTRY_FILENAME
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["КОНТРАГЕНТ", "ИНН", "№претензии", "Дата претензии", "Сумма долга"])
        # Нумерация со сбросом: старые 60, 78, затем сброс до 48, последняя — 56.
        sheet.append(["ООО «Старое-1»", None, "№60 от 01.10.2024", None, 100000])
        sheet.append(["ООО «Старое-2»", None, "№78 от 20.11.2024", None, 200000])
        sheet.append(["ООО «Новое-1»", None, "№48 от 08.04.2026", None, 300000])
        sheet.append(["ООО «Новое-2»", None, "№56 от 15.04.2026", None, 400000])
        workbook.save(self.registry_path)

    def test_resolve_registry_path_from_folder(self):
        self.assertEqual(resolve_registry_path(self.temp_dir), self.registry_path)

    def test_resolve_registry_path_missing_raises(self):
        empty_dir = tempfile.mkdtemp()
        with self.assertRaises(FileNotFoundError):
            resolve_registry_path(empty_dir)

    def test_resolve_registry_path_empty_falls_back_to_default(self):
        """Пустой путь — не ошибка: берётся папка по умолчанию из конфига."""
        expected = Path(CLAIM_REGISTRY_DEFAULT_FOLDER) / CLAIM_REGISTRY_FILENAME
        with mock.patch.object(Path, "is_file", return_value=False), \
             mock.patch.object(Path, "exists", return_value=True):
            self.assertEqual(resolve_registry_path(""), expected)
            self.assertEqual(resolve_registry_path(None), expected)

    def test_find_claim_by_inn_reads_date_from_number_cell(self):
        """Историческая строка: даты в столбце 4 нет — берётся из «№55 от 15.04.2026»."""
        self._append(["ООО «Старый»", "7712345678", "№55 от 15.04.2026", None, 1000])
        self.assertEqual(
            find_claim_by_inn("7712345678", self.temp_dir),
            {"number": "55", "date": "2026-04-15"},
        )

    def test_find_claim_by_inn_reads_separate_date_column(self):
        """Строка, созданная программой: номер в столбце 3, дата — в столбце 4."""
        self._append(["ООО «Новый»", "7701234567", "№57", "13.07.2026", 2000])
        self.assertEqual(
            find_claim_by_inn("7701234567", self.temp_dir),
            {"number": "57", "date": "2026-07-13"},
        )

    def test_find_claim_by_inn_handles_numeric_cells_and_dirty_inn(self):
        """ИНН может лежать числом, а на входе — с пробелами/дефисами."""
        self._append(["ООО «Числом»", 7799999999, 60, None, 3000])
        self.assertEqual(
            find_claim_by_inn("77-999 999 99", self.temp_dir),
            {"number": "60", "date": ""},
        )

    def test_find_claim_by_inn_returns_last_match(self):
        """У должника несколько претензий — берётся последняя (она предшествует иску)."""
        self._append(["ООО «Дубль»", "7712345678", "№55 от 15.04.2026", None, 1000])
        self._append(["ООО «Дубль»", "7712345678", "№99", "01.02.2027", 5000])
        self.assertEqual(
            find_claim_by_inn("7712345678", self.temp_dir),
            {"number": "99", "date": "2027-02-01"},
        )

    def test_find_claim_by_inn_ignores_rows_without_inn(self):
        """Исторические строки без ИНН (и строка-дубль заголовка) не находятся."""
        self._append(["ООО «Без ИНН»", None, "№58 от 01.01.2026", None, 1000])
        self._append(["КОНТРАГЕНТ", "ИНН", "№ претензии", "Дата претензии", None])
        self.assertIsNone(find_claim_by_inn("7712345678", self.temp_dir))
        self.assertIsNone(find_claim_by_inn("", self.temp_dir))

    def _append(self, row):
        """Дописывает строку в тестовый реестр."""
        workbook = openpyxl.load_workbook(self.registry_path)
        workbook.active.append(row)
        workbook.save(self.registry_path)

    def test_next_number_follows_last_row_not_global_max(self):
        """Следующий номер = последняя строка (56) + 1, а не глобальный макс (78)."""
        self.assertEqual(peek_next_claim_number(self.registry_path), 57)

    def test_append_registry_row_writes_first_empty_row(self):
        row_index = append_registry_row(
            self.registry_path, name="ООО «Ромашка»", inn="7701234567",
            number_cell="№57", claim_date=date(2026, 7, 13), amount=Decimal("269789.80"),
        )
        self.assertEqual(row_index, 6)  # 1 заголовок + 4 строки + новая
        workbook = openpyxl.load_workbook(self.registry_path, data_only=True)
        sheet = workbook.active
        written = [c.value for c in sheet[row_index]][:5]
        self.assertEqual(written[0], "ООО «Ромашка»")
        self.assertEqual(written[1], "7701234567")
        self.assertEqual(written[2], "№57")
        self.assertEqual(written[3], "13.07.2026")
        self.assertAlmostEqual(float(written[4]), 269789.80, places=2)


class InvoicesSummaryTests(unittest.TestCase):
    """Проверки подсчёта итогов из загруженного Excel со счетами."""

    def _make_invoices(self):
        temp_dir = tempfile.mkdtemp()
        path = os.path.join(temp_dir, "Счета.xlsx")
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Дата", "Номер", "Договор", "Сумма", "Срок"])
        sheet.append([date(2024, 6, 24), "65625", "2706/23-АР от 27.06.2023", 40480, date(2024, 7, 8)])
        sheet.append([date(2024, 6, 25), "65626", "2706/23-АР от 27.06.2023", 120000.50, date(2024, 7, 9)])
        workbook.save(path)
        return path

    def test_summarize_sums_column_four_and_reads_contract(self):
        summary = summarize_invoices_excel(self._make_invoices())
        self.assertEqual(summary["counted"], 2)
        self.assertEqual(summary["total_raw"], "160480.5")   # str(Decimal) без хвостового нуля
        self.assertEqual(summary["total"], "160 480,50")     # то, что видит пользователь
        self.assertEqual(summary["contract_number"], "2706/23-АР")
        self.assertEqual(summary["contract_date"], "2023-06-27")


class ClaimPackageGenerationTests(unittest.TestCase):
    """Проверки полной генерации комплекта по претензии."""

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.registry_path = Path(self.temp_dir) / CLAIM_REGISTRY_FILENAME
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["КОНТРАГЕНТ", "ИНН", "№претензии", "Дата претензии", "Сумма долга"])
        sheet.append(["ООО «Прошлое»", None, "№56 от 15.04.2026", None, 400000])
        workbook.save(self.registry_path)

    def _request(self, **overrides):
        request = {
            "defendant": {"name": "ООО «Ромашка»", "inn": "7701234567",
                          "ogrn": "1027700000000", "address": "г. Москва"},
            "claim_number": "",
            "claim_date": "2026-07-13",
            "contract_number": "2706/23-АР",
            "contract_date": "2023-06-27",
            "debt_amount": "269 789,80",
            "registry_folder": self.temp_dir,
        }
        request.update(overrides)
        return request

    def test_auto_number_is_last_plus_one(self):
        result = generate_claim_package(self._request())
        self.assertEqual(result["claim_number"], "57")

    def test_manual_number_is_used(self):
        result = generate_claim_package(self._request(claim_number="№100"))
        self.assertEqual(result["claim_number"], "100")

    def test_generates_three_temp_files(self):
        result = generate_claim_package(self._request())
        for key in ("docx_path", "pdf_path", "opis_path"):
            self.assertTrue(os.path.exists(result[key]), key)

    def test_writes_registry_row(self):
        generate_claim_package(self._request())
        workbook = openpyxl.load_workbook(self.registry_path, data_only=True)
        sheet = workbook.active
        last = None
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if any(v not in (None, "") for v in row[:5]):
                last = row
        self.assertEqual(last[0], "ООО «Ромашка»")
        self.assertEqual(last[2], "№57")
        self.assertEqual(last[3], "13.07.2026")

    def test_missing_debt_amount_raises(self):
        with self.assertRaises(ValueError):
            generate_claim_package(self._request(debt_amount=""))

    def test_export_copies_document_and_opis(self):
        result = generate_claim_package(self._request())
        out_dir = tempfile.mkdtemp()
        target = os.path.join(out_dir, "Претензия.pdf")
        export = export_claim_document({
            "source_path": result["pdf_path"],
            "target_path": target,
            "opis_path": result["opis_path"],
            "safe_name": result["safe_name"],
        })
        self.assertTrue(os.path.exists(target))
        self.assertEqual(len(export["saved"]), 2)
        self.assertTrue(any("Опись" in p for p in export["saved"]))


if __name__ == "__main__":
    unittest.main()
