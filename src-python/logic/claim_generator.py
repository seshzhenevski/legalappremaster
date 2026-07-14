# -*- coding: utf-8 -*-
"""
Слой логики: генерация досудебной претензии.

Собирает контекст письма из реквизитов должника, договора и суммы долга,
формирует три файла (претензия DOCX, претензия PDF с факсимиле, опись DOCX)
во временную папку и дописывает строку в реестр «ОТПРАВКИ ПРЕТЕНЗИЙ.xlsx».
Скачивание готовых файлов (копирование в выбранное пользователем место)
вынесено в export_claim_document, потому что sidecar живёт один запрос —
пути к временным файлам возвращаются интерфейсу и приходят обратно на шаге
скачивания.

Ничего не знает об интерфейсе, кроме функции report_progress (как и
lawsuit_generator / document_sorter_service).
"""
from __future__ import annotations

import re
import shutil
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import List, Optional

from legal_tools.config import CLAIM_REGISTRY_FILENAME, CLAIM_REGISTRY_DEFAULT_FOLDER
from legal_tools.core.formatting import (
    parse_amount_cell, parse_contract_ref, parse_date_cell, sanitize_filename, fmt,
)
from legal_tools.generators.claim_text import format_claim_number
from legal_tools.generators.claim_docx import generate_claim_docx
from legal_tools.generators.claim_pdf import generate_claim_pdf
from legal_tools.generators.lawsuit_docx import generate_opis_docx

try:
    import openpyxl
    _XL = True
except Exception:
    _XL = False

BANNER = "=" * 50
_TEMP_SUBDIR = "trivio_claim"


def parse_iso_date(iso_date_string: Optional[str]) -> Optional[date]:
    """
    Преобразует строку ISO (ГГГГ-ММ-ДД) в date. Пустое значение → None.

    Фронтенд передаёт даты полей в ISO-формате (скрытый инпут календаря);
    пустая строка означает «дата не задана».
    """
    text = (iso_date_string or "").strip()
    if not text:
        return None
    year, month, day = (int(part) for part in text.split("-"))
    return date(year, month, day)


def resolve_registry_path(registry_folder: Optional[str] = None) -> Path:
    """
    Находит файл реестра «ОТПРАВКИ ПРЕТЕНЗИЙ.xlsx» по указанному пути.

    Принимает путь к папке (обычный случай) либо прямой путь к самому
    xlsx-файлу. Если путь не задан — берёт папку по умолчанию из конфига
    (общая сетевая папка юридического департамента). Возвращает путь к
    существующему файлу или бросает понятную ошибку, если файл не найден
    (в том числе когда сетевая папка недоступна).
    """
    raw = (registry_folder or "").strip() or CLAIM_REGISTRY_DEFAULT_FOLDER
    candidate = Path(raw)
    if candidate.is_file():
        return candidate
    registry_path = candidate / CLAIM_REGISTRY_FILENAME
    if not registry_path.exists():
        raise FileNotFoundError(
            f"В папке не найден файл «{CLAIM_REGISTRY_FILENAME}»:\n{raw}"
        )
    return registry_path


def _require_openpyxl() -> None:
    if not _XL:
        raise RuntimeError(
            "Библиотека openpyxl не установлена.\n\nВыполните: pip install openpyxl"
        )


def summarize_invoices_excel(excel_path: str) -> dict:
    """
    Считает итоги из загруженного Excel со списком счетов.

    Формат совпадает с реестром для иска: столбец 3 — договор «номер от дата»,
    столбец 4 — сумма. Суммирует все распознанные суммы столбца 4 и берёт
    реквизиты договора из первой непустой ячейки столбца 3. Возвращает словарь
    с общей суммой (строкой), номером договора и датой договора в ISO — готово
    для подстановки в поля формы.
    """
    _require_openpyxl()
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    worksheet = workbook.active

    total = Decimal("0")
    counted = 0
    contract_number = ""
    contract_date_iso = ""
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        cells = (list(row) + [None] * 4)[:4]
        contract_cell, amount_cell = cells[2], cells[3]
        amount = parse_amount_cell(amount_cell)
        if amount is not None:
            total += amount
            counted += 1
        if not contract_number and contract_cell not in (None, ""):
            _full, number, date_text = parse_contract_ref(str(contract_cell))
            contract_number = number
            contract_date_iso = _contract_date_to_iso(date_text)

    if counted == 0:
        raise ValueError(
            "В выбранном файле не найдено ни одной суммы в столбце 4."
        )

    return {
        "total": fmt(total),
        "total_raw": str(total),
        "counted": counted,
        "contract_number": contract_number,
        "contract_date": contract_date_iso,
    }


def _contract_date_to_iso(date_text: str) -> str:
    """«27.06.2023» → «2023-06-27». Пустая/нераспознанная строка → «»."""
    text = (date_text or "").strip()
    match = re.match(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$", text)
    if not match:
        return ""
    day, month, year = (int(part) for part in match.groups())
    if year < 100:
        year += 2000
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return ""


def _iter_data_rows(worksheet):
    """Отдаёт (индекс_строки, значения) по строкам данных (со 2-й)."""
    for index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        yield index, row


def peek_next_claim_number(registry_path: Path) -> int:
    """
    Определяет следующий порядковый номер претензии по реестру.

    Берёт число из ПОСЛЕДНЕЙ заполненной строки столбца 3 и прибавляет 1
    (нумерация в файле периодически сбрасывается, поэтому именно последняя
    строка, а не глобальный максимум). Если подходящих строк нет — 1.
    """
    _require_openpyxl()
    workbook = openpyxl.load_workbook(registry_path, data_only=True)
    worksheet = workbook.active
    last_number = 0
    for _index, row in _iter_data_rows(worksheet):
        cell = row[2] if len(row) > 2 else None
        if cell in (None, ""):
            continue
        match = re.search(r"\d+", str(cell))
        if match:
            last_number = int(match.group())
    return last_number + 1


# Дата внутри «старого» формата столбца 3: «№55 от 15.04.2026».
_DATE_IN_NUMBER_RE = re.compile(r"от\s+(\d{1,2}\.\d{1,2}\.\d{2,4})", re.IGNORECASE)


def _cell_digits(cell) -> str:
    """Оставляет от значения ячейки только цифры (ИНН может лежать и числом)."""
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        cell = int(cell)
    return re.sub(r"\D", "", str(cell))


def _registry_number(number_cell) -> str:
    """
    Достаёт номер претензии из столбца 3 в любом из встречающихся форматов.

    В файле исторически лежат разные варианты: число (60), «№55 от 15.04.2026»
    и новый «№57». Во всех случаях возвращает голый номер («60», «55», «57»).
    """
    if number_cell in (None, ""):
        return ""
    match = re.search(r"\d+", str(number_cell))
    return match.group() if match else ""


def _registry_date_iso(date_cell, number_cell) -> str:
    """
    Достаёт дату претензии в ISO из столбца 4, а если он пуст — из столбца 3.

    В новых строках дата лежит в столбце 4, в старых — зашита в текст номера
    («№55 от 15.04.2026»). Возвращает «» , если дату распознать не удалось.
    """
    parsed = parse_date_cell(date_cell)
    if parsed is not None:
        return parsed.isoformat()
    if number_cell not in (None, ""):
        match = _DATE_IN_NUMBER_RE.search(str(number_cell))
        if match:
            parsed = parse_date_cell(match.group(1))
            if parsed is not None:
                return parsed.isoformat()
    return ""


def find_claim_by_inn(inn: str, registry_folder: Optional[str] = None) -> Optional[dict]:
    """
    Ищет в реестре последнюю претензию должника по ИНН.

    Сравнивает ИНН со столбцом 2 реестра (по цифрам, чтобы не зависеть от того,
    записан он текстом или числом) и возвращает номер и дату (ISO) ПОСЛЕДНЕЙ
    подходящей строки — именно эта претензия предшествует иску. Возвращает None,
    если строк с таким ИНН нет.

    Внимание: в исторических строках столбец ИНН не заполнялся, поэтому находятся
    только те претензии, у которых ИНН проставлен (сформированные программой либо
    заполненные вручную).
    """
    digits = _cell_digits(inn)
    if not digits:
        return None

    _require_openpyxl()
    registry_path = resolve_registry_path(registry_folder)
    workbook = openpyxl.load_workbook(registry_path, data_only=True)
    worksheet = workbook.active

    found: Optional[dict] = None
    for _index, row in _iter_data_rows(worksheet):
        inn_cell = row[1] if len(row) > 1 else None
        if _cell_digits(inn_cell) != digits:
            continue
        number_cell = row[2] if len(row) > 2 else None
        date_cell = row[3] if len(row) > 3 else None
        found = {
            "number": _registry_number(number_cell),
            "date": _registry_date_iso(date_cell, number_cell),
        }
    return found


def append_registry_row(
    registry_path: Path, *, name: str, inn: str, number_cell: str,
    claim_date: date, amount: Decimal,
) -> int:
    """
    Дописывает строку по претензии в первую пустую строку реестра.

    Столбцы: 1 — наименование должника, 2 — ИНН, 3 — номер претензии,
    4 — дата претензии (ДД.ММ.ГГГГ), 5 — сумма долга (число). Столбец 6
    (идентификатор отправления) не трогается. Возвращает индекс строки.
    Бросает понятную ошибку, если файл открыт в Excel.
    """
    _require_openpyxl()
    try:
        workbook = openpyxl.load_workbook(registry_path)
    except PermissionError:
        raise PermissionError(_locked_message(registry_path.name))
    worksheet = workbook.active

    last_data_row = 1
    for index, row in _iter_data_rows(worksheet):
        if any(value not in (None, "") for value in row[:5]):
            last_data_row = index
    target_row = last_data_row + 1

    worksheet.cell(row=target_row, column=1, value=name)
    worksheet.cell(row=target_row, column=2, value=inn or None)
    worksheet.cell(row=target_row, column=3, value=number_cell)
    worksheet.cell(row=target_row, column=4, value=claim_date.strftime("%d.%m.%Y"))
    worksheet.cell(row=target_row, column=5, value=float(amount))

    try:
        workbook.save(registry_path)
    except PermissionError:
        raise PermissionError(_locked_message(registry_path.name))
    return target_row


def _locked_message(filename: str) -> str:
    return (
        f"Не удалось записать в файл:\n{filename}\n\n"
        "Скорее всего, он открыт в Excel. Закройте файл и повторите."
    )


def _prepare_temp_dir() -> Path:
    """Создаёт (пересоздаёт) временную папку для файлов текущей претензии."""
    temp_dir = Path(tempfile.gettempdir()) / _TEMP_SUBDIR
    if temp_dir.exists():
        shutil.rmtree(temp_dir, ignore_errors=True)
    temp_dir.mkdir(parents=True, exist_ok=True)
    return temp_dir


def generate_claim_package(request: dict, report_progress=None) -> dict:
    """
    Генерирует комплект по досудебной претензии и дописывает реестр.

    Принимает запрос (реквизиты должника, номер/дата претензии, номер/дата
    договора, сумма долга, папка с реестром) и необязательную функцию
    report_progress. Создаёт претензию DOCX и PDF (с факсимиле) и опись во
    временной папке, затем дописывает строку в «ОТПРАВКИ ПРЕТЕНЗИЙ.xlsx».
    Возвращает пути к временным файлам и сводку для интерфейса; сами файлы
    пользователь скачивает через export_claim_document.
    """
    if report_progress is None:
        report_progress = lambda message, percent=None: None

    defendant = request.get("defendant", {})
    name = (defendant.get("name") or "").strip()
    if not name:
        raise ValueError("Укажите наименование должника.")

    claim_date = parse_iso_date(request.get("claim_date")) or date.today()
    debt_amount = parse_amount_cell(request.get("debt_amount"))
    if debt_amount is None or debt_amount <= 0:
        raise ValueError("Укажите сумму долга больше нуля (или загрузите Excel со счетами).")

    contract_number = (request.get("contract_number") or "").strip()
    contract_date = parse_iso_date(request.get("contract_date"))
    registry_path = resolve_registry_path(request.get("registry_folder"))

    report_progress(BANNER)
    report_progress("▶ ФОРМИРОВАНИЕ ДОСУДЕБНОЙ ПРЕТЕНЗИИ")
    report_progress(BANNER)
    report_progress("")

    manual_number = format_claim_number(request.get("claim_number", ""))
    if manual_number:
        claim_number = manual_number
        report_progress(f"📌 Номер претензии (введён вручную): № {claim_number}", percent=10)
    else:
        next_number = peek_next_claim_number(registry_path)
        claim_number = str(next_number)
        report_progress(
            f"📌 Номер претензии сформирован автоматически: № {claim_number}",
            percent=10,
        )
    number_cell = f"№{claim_number}"

    context = {
        "defendant_name": name,
        "defendant_address": (defendant.get("address") or "").strip(),
        "claim_number": claim_number,
        "claim_date": claim_date,
        "contract_number": contract_number,
        "contract_date": contract_date,
        "debt_amount": debt_amount,
    }
    report_progress(f"Должник: {name}")
    report_progress(f"Сумма долга: {fmt(debt_amount)} руб.")
    if contract_number:
        report_progress(f"Договор: № {contract_number}"
                        + (f" от {contract_date:%d.%m.%Y}" if contract_date else ""))
    report_progress("")

    temp_dir = _prepare_temp_dir()
    safe_name = sanitize_filename(name, "Должник")
    docx_path = temp_dir / f"Досудебная претензия ({safe_name}).docx"
    pdf_path = temp_dir / f"Досудебная претензия ({safe_name}).pdf"
    opis_path = temp_dir / f"Опись ({safe_name}).docx"

    report_progress("📄 Формирование претензии (DOCX)...", percent=35)
    generate_claim_docx(str(docx_path), context)
    report_progress(f"✓ {docx_path.name}")

    report_progress("🖊 Формирование претензии (PDF с факсимиле)...", percent=60)
    generate_claim_pdf(str(pdf_path), context, with_facsimile=True)
    report_progress(f"✓ {pdf_path.name}")

    report_progress("📋 Формирование описи почтового вложения...", percent=80)
    generate_opis_docx(str(opis_path), name, document_title="Досудебная претензия")
    report_progress(f"✓ {opis_path.name}")
    report_progress("")

    report_progress("🧾 Запись в реестр «ОТПРАВКИ ПРЕТЕНЗИЙ.xlsx»...", percent=92)
    row_index = append_registry_row(
        registry_path, name=name, inn=(defendant.get("inn") or "").strip(),
        number_cell=number_cell, claim_date=claim_date, amount=debt_amount,
    )
    report_progress(f"✓ Строка {row_index}: № {claim_number}, {claim_date:%d.%m.%Y}, "
                    f"{fmt(debt_amount)} руб.")
    report_progress("")

    report_progress(BANNER, percent=100)
    report_progress("✅ ГОТОВО — можно скачивать DOCX или PDF")
    report_progress(BANNER)

    return {
        "docx_path": str(docx_path),
        "pdf_path": str(pdf_path),
        "opis_path": str(opis_path),
        "safe_name": safe_name,
        "claim_number": claim_number,
        "claim_date": claim_date.strftime("%d.%m.%Y"),
        "total_debt": fmt(debt_amount),
        "registry_row": row_index,
    }


def export_claim_document(params: dict) -> dict:
    """
    Копирует готовый документ претензии в выбранное пользователем место.

    Принимает путь к временному файлу претензии (source_path), путь
    сохранения (target_path), путь к временной описи (opis_path) и
    безопасное имя должника (safe_name). Кладёт претензию по target_path, а
    опись — рядом (в ту же папку, всегда в .docx). Возвращает список
    сохранённых файлов. Понятно сообщает, если файл занят.
    """
    source_path = Path(params["source_path"])
    target_path = Path(params["target_path"])
    safe_name = params.get("safe_name") or "Должник"

    saved: List[str] = []
    try:
        shutil.copyfile(source_path, target_path)
        saved.append(str(target_path))
        opis_source = params.get("opis_path")
        if opis_source and Path(opis_source).exists():
            opis_target = target_path.parent / f"Опись ({safe_name}).docx"
            shutil.copyfile(opis_source, opis_target)
            saved.append(str(opis_target))
    except PermissionError:
        raise PermissionError(_locked_message(target_path.name))
    return {"saved": saved}
